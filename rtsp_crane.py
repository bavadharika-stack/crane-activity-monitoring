"""
rtsp_crane_monitor.py
─────────────────────────────────────────────────────────────────────────────
Connects to an RTSP camera stream, runs your trained YOLOv8 crane-action
model on every frame LIVE, draws bounding-box annotations on the video,
and maintains a live-updating Excel report logging:

    Start Time | End Time | Duration (s) | Action

...every time the detected crane action changes (e.g. idle -> bucket_loaded_lifted).

SETUP (run once):
    pip install ultralytics opencv-python openpyxl

USAGE:
    python rtsp_crane_monitor.py --model "C:/Users/Aravind SJ/Downloads/runs/detect/training_runs/crane_v1-6/weights/best.pt" --rtsp "rtsp://user:pass@camera-ip:554/stream1"

    Optional flags:
    --conf 0.5            minimum confidence to accept a detection (default 0.5)
    --debounce 5            consecutive frames required before accepting a state change (default 5)
    --output "report.xlsx"  path to the Excel report file (default ./crane_report.xlsx)
    --no-display              run headless (no video window) - useful for background/server use
    --resize 960               resize frame width for display/inference speed (height auto-scaled)
    --save-video "out.mp4"      path to save the annotated live recording (default ./crane_recording.mp4, set to '' to disable)
    --video-fps 15              FPS to record the saved video at (default 15)
    --box-thickness 4            bounding box / label line thickness (default 4, increase for bigger boxes/text)

CONTROLS while running:
    Press 'q' in the video window to quit cleanly (this also closes the final
    open segment and saves the Excel file).
─────────────────────────────────────────────────────────────────────────────
"""

import argparse
import os
import sys
import time
from datetime import datetime

import cv2

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.utils import get_column_letter
except ImportError:
    print("ERROR: openpyxl not installed. Run: pip install openpyxl")
    sys.exit(1)

try:
    from ultralytics import YOLO
except ImportError:
    print("ERROR: ultralytics not installed. Run: pip install ultralytics")
    sys.exit(1)


# ─────────────────────────────────────────────────────────────────────────────
# EXCEL REPORT HANDLING
# ─────────────────────────────────────────────────────────────────────────────

REPORT_HEADERS = ["Start Time", "End Time", "Duration (s)", "Action"]


def init_report(path: str):
    """Create a new workbook with headers, or load an existing one to append to."""
    if os.path.isfile(path):
        try:
            wb = load_workbook(path)
            ws = wb.active
            print(f"📄 Found existing report, appending to it: {path}")
            return wb, ws
        except Exception as e:
            print(f"⚠️  Could not open existing report ({e}). Creating a new one.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Crane Action Log"
    ws.append(REPORT_HEADERS)
    # widen columns a bit for readability
    widths = [22, 22, 14, 28]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    wb.save(path)
    print(f"📄 Created new report: {path}")
    return wb, ws


def append_segment(wb, ws, path: str, start_dt: datetime, end_dt: datetime, action: str):
    duration_s = round((end_dt - start_dt).total_seconds(), 1)
    ws.append([
        start_dt.strftime("%Y-%m-%d %H:%M:%S"),
        end_dt.strftime("%Y-%m-%d %H:%M:%S"),
        duration_s,
        action,
    ])
    try:
        wb.save(path)
    except PermissionError:
        print(f"⚠️  Could not save {path} — close it in Excel if you have it open! "
              "Segment was logged in memory and will retry on the next update.")
    return duration_s


# ─────────────────────────────────────────────────────────────────────────────
# MAIN MONITORING LOOP
# ─────────────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(description="Live RTSP crane-action monitor with Excel reporting.")
    parser.add_argument("--model", required=True, help="Path to your trained best.pt weights")
    parser.add_argument("--rtsp", required=True, help="RTSP stream URL")
    parser.add_argument("--conf", type=float, default=0.5, help="Minimum confidence threshold")
    parser.add_argument("--debounce", type=int, default=5,
                         help="Consecutive frames a new state must persist before being accepted")
    parser.add_argument("--output", default="./crane_report.xlsx", help="Path to Excel report")
    parser.add_argument("--no-display", action="store_true", help="Run without showing the video window")
    parser.add_argument("--resize", type=int, default=960, help="Resize frame width for inference/display")
    parser.add_argument("--save-video", default="./crane_recording.mp4",
                         help="Path to save the annotated live recording (set to '' to disable saving)")
    parser.add_argument("--video-fps", type=float, default=15.0,
                         help="FPS to record the output video at (default 15)")
    parser.add_argument("--box-thickness", type=int, default=4,
                         help="Bounding box / label line thickness (default 4, bigger = thicker boxes & text)")
    args = parser.parse_args()

    # ── Load model ─────────────────────────────────────────────────────────
    if not os.path.isfile(args.model):
        print(f"❌ ERROR: model file not found: {args.model}")
        sys.exit(1)

    print(f"Loading model: {args.model}")
    model = YOLO(args.model)
    class_names = model.names  # dict: id -> name
    print(f"Classes: {list(class_names.values())}\n")

    # ── Connect to RTSP stream ────────────────────────────────────────────
    print(f"Connecting to RTSP stream: {args.rtsp}")
    # FFMPEG backend + low-latency hint, helps a lot with flaky RTSP feeds on Windows
    cap = cv2.VideoCapture(args.rtsp, cv2.CAP_FFMPEG)
    cap.set(cv2.CAP_PROP_BUFFERSIZE, 1)  # minimize latency / stale-frame buildup

    if not cap.isOpened():
        print("❌ ERROR: Could not open RTSP stream. Check the URL, credentials, and network path.")
        sys.exit(1)

    print("✅ Stream connected. Starting live monitoring... (press 'q' in the video window to quit)\n")

    # ── Excel report setup ────────────────────────────────────────────────
    wb, ws = init_report(args.output)

    # ── Video recording setup ──────────────────────────────────────────────
    video_writer = None  # created lazily once we know the actual frame size

    # ── State tracking ────────────────────────────────────────────────────
    current_action = None          # the currently "locked in" / accepted state
    current_action_start = None    # datetime when current_action started
    candidate_action = None        # a possible new state being debounced
    candidate_count = 0            # how many consecutive frames candidate_action has appeared

    NO_DETECTION_LABEL = "no_detection"  # used when nothing is detected above conf threshold

    frame_count = 0
    fps_timer = time.time()
    fps_counter = 0
    display_fps = 0.0

    consecutive_read_failures = 0
    MAX_READ_FAILURES = 30  # ~ a few seconds of dropped frames before we bail

    try:
        while True:
            ret, frame = cap.read()

            if not ret or frame is None:
                consecutive_read_failures += 1
                print(f"⚠️  Frame read failed ({consecutive_read_failures}/{MAX_READ_FAILURES})")
                if consecutive_read_failures >= MAX_READ_FAILURES:
                    print("❌ Too many consecutive read failures — RTSP stream appears down. Exiting.")
                    break
                time.sleep(0.2)
                continue
            consecutive_read_failures = 0

            frame_count += 1

            # Resize for consistent inference speed / display
            if args.resize:
                h, w = frame.shape[:2]
                scale = args.resize / w
                frame = cv2.resize(frame, (args.resize, int(h * scale)))

            # ── Run inference ──────────────────────────────────────────────
            results = model.predict(frame, conf=args.conf, verbose=False)
            result = results[0]
            annotated_frame = result.plot(line_width=args.box_thickness)  # frame with boxes/labels drawn

            # Create the video writer once we know the actual annotated frame size
            if args.save_video and video_writer is None:
                h, w = annotated_frame.shape[:2]
                fourcc = cv2.VideoWriter_fourcc(*"mp4v")
                video_writer = cv2.VideoWriter(args.save_video, fourcc, args.video_fps, (w, h))
                if video_writer.isOpened():
                    print(f"🎥 Recording annotated video to: {os.path.abspath(args.save_video)}")
                else:
                    print(f"⚠️  Could not open video writer for {args.save_video} — recording disabled.")
                    video_writer = None

            # ── Determine the dominant detected action this frame ─────────
            detected_label = NO_DETECTION_LABEL
            if len(result.boxes) > 0:
                # Pick the highest-confidence detection as "the" current action.
                # (If you'd rather use majority vote across all boxes, swap this
                #  block for a Counter().most_common(1) on all class ids.)
                confs = result.boxes.conf.cpu().numpy()
                cls_ids = result.boxes.cls.cpu().numpy().astype(int)
                best_idx = confs.argmax()
                detected_label = class_names[cls_ids[best_idx]]

            now = datetime.now()

            # ── State machine with debounce ────────────────────────────────
            if current_action is None:
                # very first observed state — start tracking immediately
                current_action = detected_label
                current_action_start = now
                candidate_action = None
                candidate_count = 0
            elif detected_label == current_action:
                # still the same action — reset any pending candidate
                candidate_action = None
                candidate_count = 0
            else:
                # detected something different from the locked-in state
                if detected_label == candidate_action:
                    candidate_count += 1
                else:
                    candidate_action = detected_label
                    candidate_count = 1

                if candidate_count >= args.debounce:
                    # candidate has persisted long enough -> accept the state change
                    segment_end = now
                    duration_s = append_segment(
                        wb, ws, args.output,
                        current_action_start, segment_end, current_action
                    )
                    print(f"📝 Logged: {current_action:28s} "
                          f"{current_action_start.strftime('%H:%M:%S')} -> "
                          f"{segment_end.strftime('%H:%M:%S')}  ({duration_s}s)")

                    current_action = candidate_action
                    current_action_start = segment_end
                    candidate_action = None
                    candidate_count = 0

            # ── Overlay live status on the video ───────────────────────────
            elapsed = (now - current_action_start).total_seconds()
            overlay_text = f"Action: {current_action}   Duration: {elapsed:.1f}s"
            cv2.rectangle(annotated_frame, (0, 0), (annotated_frame.shape[1], 40), (0, 0, 0), -1)
            cv2.putText(annotated_frame, overlay_text, (10, 28),
                        cv2.FONT_HERSHEY_SIMPLEX, 0.7, (0, 255, 255), 2)

            # FPS counter (helps you judge if inference is keeping up with the stream)
            fps_counter += 1
            if time.time() - fps_timer >= 1.0:
                display_fps = fps_counter / (time.time() - fps_timer)
                fps_counter = 0
                fps_timer = time.time()
            cv2.putText(annotated_frame, f"FPS: {display_fps:.1f}",
                        (annotated_frame.shape[1] - 140, 28),
                        cv2.FONT_HERSHEY_SIMPLEX, 0.6, (0, 255, 0), 2)

            # ── Write this frame to the recording ──────────────────────────
            if video_writer is not None:
                video_writer.write(annotated_frame)

            # ── Display ─────────────────────────────────────────────────
            if not args.no_display:
                cv2.imshow("Crane Live Monitor (press 'q' to quit)", annotated_frame)
                if cv2.waitKey(1) & 0xFF == ord('q'):
                    print("\n🛑 'q' pressed — stopping...")
                    break

    except KeyboardInterrupt:
        print("\n🛑 Interrupted by user (Ctrl+C)...")

    finally:
        # Close out the final open segment so it's not lost
        if current_action is not None and current_action_start is not None:
            segment_end = datetime.now()
            duration_s = append_segment(
                wb, ws, args.output,
                current_action_start, segment_end, current_action
            )
            print(f"📝 Final segment logged: {current_action:28s} "
                  f"{current_action_start.strftime('%H:%M:%S')} -> "
                  f"{segment_end.strftime('%H:%M:%S')}  ({duration_s}s)")

        cap.release()
        if video_writer is not None:
            video_writer.release()
            print(f"🎥 Recording saved at: {os.path.abspath(args.save_video)}")
        cv2.destroyAllWindows()
        print(f"\n✅ Report saved at: {os.path.abspath(args.output)}")
        print("Done.")


if __name__ == "__main__":
    main()
